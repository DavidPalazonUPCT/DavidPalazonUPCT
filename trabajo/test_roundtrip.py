import sys; sys.path.insert(0,'trabajo')
from docx import Document
from docx_builder import body_para, body_para_runs, h1, h2, bib_para, empty_para

doc = Document("insumos/Plan_Investigacion_v36_APA7.docx")
body = doc.element.body
sectPr = body.find(qn('w:sectPr')) if (qn:=__import__('docx.oxml.ns',fromlist=['qn']).qn) else None
# find final sectPr (direct child)
from docx.oxml.ns import qn
sectPr = body.findall(qn('w:sectPr'))[-1]

# Remove all children after cover (keep paras 0..14 + their sectPr is inside p14). Cover end = para index 14.
paras = doc.paragraphs
cover_end_p = paras[14]._p  # 'Murcia...' holds embedded sectPr
# Collect children to remove: everything after cover_end_p except final sectPr
removing = False
to_remove = []
for child in list(body.iterchildren()):
    if child is cover_end_p:
        removing = True
        continue
    if removing:
        if child is sectPr:
            break
        to_remove.append(child)
for c in to_remove:
    body.remove(c)

# Insert test content before sectPr
def ins(el): sectPr.addprevious(el)
ins(h1("1. Sección de prueba"))
ins(body_para("Este es un párrafo de cuerpo justificado con sangría de primera línea para comprobar el formato Times New Roman 12 puntos."))
ins(h2("1.1. Subsección de prueba"))
ins(body_para_runs([("Palabras clave: ", True, True), ("uno; dos; tres.", False, False)]))
ins(h1("2. Bibliografía de prueba"))
ins(bib_para("Autor, A. (2024). Un título de referencia con sangría francesa para verificar el formato APA. Revista, 1(1), 1-10."))

doc.save("trabajo/test_build.docx")
print("saved test_build.docx")
# reload to validate
d2 = Document("trabajo/test_build.docx")
print("reloaded OK, paragraphs:", len(d2.paragraphs))
for p in d2.paragraphs[-6:]:
    print(f"  <{p.style.name}> {p.text[:45]!r}")
